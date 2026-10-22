"""Unit tests for parsing, chunking and fusion (no network, no models)."""
from __future__ import annotations

import json
from collections import Counter
from pathlib import Path

import pytest

from app.rag.blocks import Block
from app.rag.chunker import TARGET_CHARS, chunk_blocks, split_sentences
from app.rag.formats import sniff_ok
from app.rag.parsers import parse_file
from app.rag.retriever import focus_tables, rrf, tokenize
from tests.test_rag_fixtures import (
    CSV_HEADER, DOCX_SENTENCES, PDF_SENTENCES, make_csv, make_docx, make_pdf, make_scanned_pdf, normalise,
)


def _all_present(sentences: list[str], texts: list[str]) -> list[str]:
    joined = [normalise(t) for t in texts]
    return [s for s in sentences if not any(normalise(s) in t for t in joined)]


def test_sentence_split_keeps_every_character() -> None:
    text = 'He said "hi." Then he left. See Fig. 3 for details. Dr. Rao agreed! e.g. this stays.\nNew line'
    parts = split_sentences(text)
    assert "".join(parts).replace(" ", "") == text.replace(" ", "").replace("\n", "")
    assert 'He said "hi."' in parts
    assert any(p.startswith("See Fig. 3") for p in parts)


def test_chunker_never_drops_text_and_respects_limits() -> None:
    sentences = [f"Sentence number {i} talks about topic {i % 7} in some detail here." for i in range(400)]
    blocks = [Block(" ".join(sentences[i:i + 25]), page=1 + i // 100, section=f"Part {i // 100}", kind="text")
              for i in range(0, 400, 25)]
    chunks = chunk_blocks(blocks)
    assert not _all_present(sentences, [c.text for c in chunks])
    assert all(len(c.text) <= 2 * TARGET_CHARS for c in chunks)
    assert sum(len(c.text) for c in chunks) / len(chunks) > TARGET_CHARS * 0.6
    # sections are hard boundaries and are prepended to the embedded text only
    for c in chunks:
        assert c.embed_text.startswith(f"Section: {c.section}\n\n") and not c.text.startswith("Section:")
        part = int(c.section.split()[-1])
        assert c.page_start == c.page_end == part + 1


def test_chunker_overlap_between_neighbours() -> None:
    text = " ".join(f"Fact {i} is that the value of item {i} equals {i * 3}." for i in range(120))
    chunks = chunk_blocks([Block(text, 1, None, "text")])
    assert len(chunks) >= 3
    for a, b in zip(chunks, chunks[1:]):
        last_sentence = split_sentences(a.text)[-1]
        assert last_sentence in b.text[: int(TARGET_CHARS * 0.2)]  # ~15% sentence overlap


def test_chunker_keeps_tables_whole_and_splits_huge_ones_with_header() -> None:
    small = "| a | b |\n|---|---|\n" + "\n".join(f"| {i} | {i * 2} |" for i in range(20))
    big = "| key | value |\n|---|---|\n" + "\n".join(f"| row{i} | {'x' * 40} |" for i in range(300))
    chunks = chunk_blocks([Block("Intro sentence.", 1, "S", "text"), Block(small, 1, "S", "table"),
                           Block(big, 2, "T", "table")])
    assert any(small in c.text for c in chunks)
    big_chunks = [c for c in chunks if c.section == "T"]
    assert len(big_chunks) > 3
    assert all(c.text.startswith("| key | value |\n|---|---|") for c in big_chunks)
    assert not _all_present([f"| row{i} |" for i in range(300)], [c.text for c in big_chunks])


def test_chunker_keeps_long_captions_wide_headers_and_long_headings() -> None:
    caption = "\n".join(f"Caption line {i} " + "x" * 80 for i in range(40))
    wide = "| " + " | ".join(f"column_{i}" for i in range(120)) + " |\n|" + "---|" * 120
    rows = "\n".join("| " + " | ".join(str(r * i) for i in range(120)) + " |" for r in range(30))
    for table in (f"{caption}\n{wide}\n{rows}", f"{caption}\n| a | b |\n|---|---|\n" + "\n".join(
            f"| {i} | {'y' * 50} |" for i in range(100)), f"{caption}\n| only | header |\n|---|---|"):
        chunks = chunk_blocks([Block(table, 1, "S", "table")])
        words = Counter(w for c in chunks for w in c.text.split())
        assert all(words[w] >= n for w, n in Counter(table.split()).items()), "table text was lost"
        assert all(len(c.text) <= 2 * TARGET_CHARS for c in chunks)
    chunks = chunk_blocks([Block("word " * 2000, 1, "H", "heading"), Block("Body.", 1, "H", "text")])
    assert all(len(c.text) <= 2 * TARGET_CHARS for c in chunks) and "Body." in chunks[-1].text


def test_wide_csv_keeps_overview_and_labels_every_row(tmp_path: Path) -> None:
    import pandas as pd

    cols: dict[str, list[int]] = {f"feature_{i:02d}_measurement": [r * i for r in range(30)] for i in range(80)}
    cols["SalePrice"] = [100000 + r for r in range(30)]
    path = tmp_path / "houses.csv"
    pd.DataFrame(cols).to_csv(path, index=False)
    chunks = chunk_blocks(parse_file(path, "houses.csv").blocks)
    text = "\n".join(c.text for c in chunks)
    assert "30 rows × 81 columns" in text and all(f"- {c} (" in text for c in cols)
    assert "SalePrice: 100016" in next(c.text for c in chunks if "Row 17:" in c.text)
    assert all(len(c.text) <= 2 * TARGET_CHARS for c in chunks)


def test_heading_only_section_is_folded_into_next() -> None:
    blocks = [Block("2 Methods", 1, "2 Methods", "heading"),
              Block("2.1 Data", 1, "2 Methods › 2.1 Data", "heading"),
              Block("We collected data.", 1, "2 Methods › 2.1 Data", "text")]
    chunks = chunk_blocks(blocks)
    assert len(chunks) == 1
    assert chunks[0].text.startswith("2 Methods") and "We collected data." in chunks[0].text


def test_pdf_parser_pages_headings_and_tables(tmp_path: Path) -> None:
    parsed = parse_file(make_pdf(tmp_path / "report.pdf"), "report.pdf")
    assert parsed.page_count == 3 and not parsed.ocr_used
    headings = [(b.text, b.page) for b in parsed.blocks if b.kind == "heading"]
    assert ("2 Photosynthesis Experiments", 2) in headings and ("3 Orbital Mechanics", 3) in headings
    tables = [b for b in parsed.blocks if b.kind == "table"]
    assert len(tables) == 1 and tables[0].page == 2 and "Maize" in tables[0].text and "|" in tables[0].text
    orbit = next(b for b in parsed.blocks if "geostationary" in b.text)
    assert orbit.page == 3 and orbit.section and orbit.section.endswith("3 Orbital Mechanics")
    assert not _all_present(PDF_SENTENCES, [b.text for b in parsed.blocks])


def test_scanned_pdf_without_ocr_is_flagged(tmp_path: Path) -> None:
    parsed = parse_file(make_scanned_pdf(tmp_path / "scan.pdf"), "scan.pdf")
    assert parsed.blocks == [] and parsed.ocr_needed


def test_docx_parser_sections_tables_pages(tmp_path: Path) -> None:
    parsed = parse_file(make_docx(tmp_path / "review.docx"), "review.docx")
    assert not _all_present(DOCX_SENTENCES, [b.text for b in parsed.blocks])
    hiring = next(b for b in parsed.blocks if "machine learning engineers" in b.text)
    assert hiring.section == "Quarterly Review › Hiring Plan" and hiring.page == 2
    table = next(b for b in parsed.blocks if b.kind == "table")
    assert "| Asia Pacific | 12.4M |" in table.text and table.page == 1
    assert parsed.page_count == 2


def test_docx_text_inside_content_controls_is_kept(tmp_path: Path) -> None:
    import docx
    from docx.oxml import OxmlElement

    d = docx.Document()
    d.add_heading("Template Title", level=1)
    paras = [d.add_paragraph(f"Body sentence number {i} inside a content control.") for i in range(3)]
    sdt, content = OxmlElement("w:sdt"), OxmlElement("w:sdtContent")
    sdt.append(content)
    paras[0]._p.addprevious(sdt)
    for p in paras:
        content.append(p._p)
    path = tmp_path / "template.docx"
    d.save(str(path))
    texts = [b.text for b in parse_file(path, "template.docx").blocks]
    assert texts[0] == "Template Title" and all(any(f"number {i} inside" in t for t in texts) for i in range(3))


def test_csv_parser_overview_and_row_groups(tmp_path: Path) -> None:
    parsed = parse_file(make_csv(tmp_path / "weather.csv"), "weather.csv")
    overview, *groups = parsed.blocks
    assert "300 rows × 4 columns" in overview.text and "temperature_c (number" in overview.text
    assert "Numeric summary" in overview.text
    header = "| " + " | ".join(CSV_HEADER) + " |"
    assert len(groups) >= 300 // 40 and all(header in g.text for g in groups)
    assert sum(g.text.count("\n| ") - 1 for g in groups) == 300  # every row exactly once (minus header line)


def test_json_and_markdown_parsers(tmp_path: Path) -> None:
    path = tmp_path / "config.json"
    path.write_text(json.dumps({"model": {"name": "bge-small", "dim": 384}, "notes": "hybrid search"}), "utf-8")
    parsed = parse_file(path, "config.json")
    assert [b.section for b in parsed.blocks] == ["model", "notes"] and '"dim": 384' in parsed.blocks[0].text
    many = tmp_path / "strings.json"
    many.write_text(json.dumps({f"greeting_{i}": f"Hello number {i}" for i in range(3000)}), "utf-8")
    blocks = parse_file(many, "strings.json").blocks
    assert 20 < len(blocks) < 300 and '"greeting_2999": "Hello number 2999"' in blocks[-1].text
    assert sum(b.text.count('"greeting_') for b in blocks) == 3000

    md = tmp_path / "notes.md"
    md.write_text("# Title\n\nIntro para.\n\n## Install\n\n```bash\npip install x\n```\n\n| a | b |\n|---|---|\n| 1 | 2 |\n",
                  "utf-8")
    blocks = parse_file(md, "notes.md").blocks
    assert [b.kind for b in blocks] == ["heading", "text", "heading", "code", "table"]
    assert blocks[3].section == "Title › Install"


def test_xlsx_sheets_and_html(tmp_path: Path) -> None:
    import pandas as pd

    xlsx = tmp_path / "budget.xlsx"
    with pd.ExcelWriter(xlsx) as writer:
        pd.DataFrame({"item": ["GPU", "CPU"], "cost": [1200, 300]}).to_excel(writer, sheet_name="Hardware", index=False)
        pd.DataFrame({"name": ["Asha"], "role": ["ML engineer"]}).to_excel(writer, sheet_name="Team", index=False)
    blocks = parse_file(xlsx, "budget.xlsx").blocks
    assert {b.section for b in blocks} == {"Sheet: Hardware", "Sheet: Team"}
    assert any("| GPU | 1200 |" in b.text for b in blocks) and any("ML engineer" in b.text for b in blocks)

    html = tmp_path / "page.html"
    html.write_text("<html><head><title>Sphura Docs</title><script>var x=1;</script></head><body><article>"
                    "<h1>Getting started</h1><p>Upload a PDF and ask questions about it with citations.</p>"
                    "<h2>Limits</h2><p>Files up to twenty five megabytes are accepted by the server.</p>"
                    "</article></body></html>", "utf-8")
    blocks = parse_file(html, "page.html").blocks
    text = " ".join(b.text for b in blocks)
    assert "twenty five megabytes" in text and "var x" not in text
    assert any(b.section and b.section.startswith("Sphura Docs") for b in blocks)


def test_gemini_ocr_path_for_scanned_pdf(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """With a key, image-only pages are rendered, sent as a data URL and parsed as markdown (HTTP mocked)."""
    import httpx

    from app.core.config import settings
    from app.rag import ocr

    sent: list[dict] = []

    def fake_post(url: str, json: dict, headers: dict, timeout: float) -> httpx.Response:
        sent.append(json)
        content = "```markdown\n# Scanned Invoice\n\nTotal amount due is 4,250 rupees.\n```"
        return httpx.Response(200, json={"choices": [{"message": {"content": content}}]},
                              request=httpx.Request("POST", url))

    monkeypatch.setattr(settings, "gemini_api_key", "test-key")
    monkeypatch.setattr(ocr.httpx, "post", fake_post)
    parsed = parse_file(make_scanned_pdf(tmp_path / "scan.pdf"), "scan.pdf")
    assert parsed.ocr_used and not parsed.ocr_needed
    assert [(b.kind, b.text, b.page) for b in parsed.blocks] == [
        ("heading", "Scanned Invoice", 1), ("text", "Total amount due is 4,250 rupees.", 1)]
    image_part = sent[0]["messages"][0]["content"][1]["image_url"]["url"]
    assert image_part.startswith("data:image/png;base64,") and sent[0]["model"] == ocr.GEMINI_OCR_MODEL


def test_magic_sniff() -> None:
    assert sniff_ok(".pdf", b"%PDF-1.7\n...")
    assert not sniff_ok(".pdf", b"hello world")
    assert sniff_ok(".docx", b"PK\x03\x04rest") and not sniff_ok(".xlsx", b"%PDF-")
    assert sniff_ok(".png", b"\x89PNG\r\n\x1a\nxxxx") and not sniff_ok(".jpg", b"\x89PNG")
    assert sniff_ok(".csv", b"a,b\n1,2") and not sniff_ok(".txt", b"MZ\x90\x00\x03")


def test_focus_tables_keeps_header_and_matching_rows() -> None:
    table = "Rows 1–20\n\n| id | city | rain |\n|---|---|---|\n" + "\n".join(
        f"| {i} | {'Zanzibar' if i == 7 else 'Pune'} | {i} |" for i in range(20))
    focused = focus_tables(table, set(tokenize("rain in Zanzibar")))
    assert focused.splitlines() == ["Rows 1–20", "", "| id | city | rain |", "|---|---|---|", "| 7 | Zanzibar | 7 |"]
    assert focus_tables("short prose", {"x"}) == "short prose"
    unmatched = focus_tables(table, {"nothing"})
    assert "| 0 | Pune | 0 |" in unmatched and "| 19 | Pune | 19 |" not in unmatched


def test_tokenizer_and_rrf() -> None:
    assert tokenize("The Tables of the Results!") == ["table", "result"]
    fused = rrf([["a", "b", "c"], ["c", "a"]])
    assert max(fused, key=fused.get) == "a" and fused["c"] > fused["b"]
