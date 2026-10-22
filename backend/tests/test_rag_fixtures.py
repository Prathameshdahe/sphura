"""Generated test documents for the RAG tests (helpers only — no tests here)."""
from __future__ import annotations

import csv
import re
from pathlib import Path

import docx
import pymupdf

INTRO = [
    "Sphura is a research assistant that reads long documents and answers questions with citations.",
    "The system was evaluated on a corpus of technical reports collected over three semesters.",
    "Every answer must point back to the page that supports it.",
]
FILLER = [
    f"Observation number {i} in the {topic} log confirmed the expected behaviour of the {thing}."
    for i, (topic, thing) in enumerate(
        [("calibration", "sensor array"), ("maintenance", "cooling loop"), ("deployment", "edge gateway"),
         ("validation", "parser module"), ("integration", "message queue"), ("stress", "storage tier"),
         ("regression", "ranking stage"), ("acceptance", "upload service"), ("security", "token issuer"),
         ("latency", "vector index"), ("throughput", "batch embedder"), ("usability", "chat window")], start=1)
]
PHOTO = [
    "The chlorophyll absorption peak was measured at 680 nanometers during the greenhouse trial.",
    "Leaves exposed to red light showed a markedly higher rate of oxygen evolution.",
    "Table 1 lists the crop yields recorded at the end of the growing season.",
]
YIELD_TABLE = [["Crop", "Yield (t/ha)", "Water (mm)"], ["Wheat", "4.2", "450"], ["Rice", "5.1", "1200"],
               ["Maize", "6.3", "600"]]
ORBIT = [
    "The relay satellite maintained a geostationary orbit at an altitude of 35786 kilometers above the equator.",
    "Station keeping manoeuvres consumed roughly two kilograms of xenon propellant per year.",
    "Ground controllers uplinked new ephemeris data every Tuesday morning.",
]
PDF_SENTENCES = INTRO + FILLER + PHOTO + ORBIT

DOCX_REVENUE = [
    "Revenue for the quarter grew by eighteen percent compared with the previous year.",
    "Most of the growth came from enterprise subscriptions in the Asia Pacific region.",
]
DOCX_HIRING = [
    "The company plans to hire forty two machine learning engineers before the end of March.",
    "Interviews will be run by a panel of three senior staff members.",
]
DOCX_SENTENCES = DOCX_REVENUE + DOCX_HIRING

CSV_HEADER = ["id", "city", "temperature_c", "rainfall_mm"]
_CITIES = ["Pune", "Delhi", "Mumbai", "Chennai", "Kolkata", "Jaipur", "Lucknow", "Bhopal", "Indore", "Surat"]


def normalise(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def _textbox(page: pymupdf.Page, rect: pymupdf.Rect, text: str, size: float = 11) -> float:
    rest = page.insert_textbox(rect, text, fontsize=size, fontname="helv")
    assert rest >= 0, "test text overflowed its box"
    return rect.y0 + (rect.height - rest) + 8


def _heading(page: pymupdf.Page, y: float, text: str, size: float = 16) -> float:
    page.insert_text((72, y + size), text, fontsize=size, fontname="hebo")
    return y + size + 12


def _table(page: pymupdf.Page, y: float, rows: list[list[str]]) -> float:
    cw, rh = 130, 22
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            rect = pymupdf.Rect(72 + c * cw, y + r * rh, 72 + (c + 1) * cw, y + (r + 1) * rh)
            page.draw_rect(rect, color=(0, 0, 0), width=0.8)
            page.insert_text((rect.x0 + 5, rect.y1 - 7), val, fontsize=10, fontname="helv")
    return y + len(rows) * rh + 14


def make_pdf(path: Path) -> Path:
    """3 pages: headings (big bold font), prose, a ruled table on page 2."""
    doc = pymupdf.open()
    width = 470
    p1 = doc.new_page()
    y = _heading(p1, 50, "Sphura Technical Report", 22)
    y = _heading(p1, y, "1 Introduction")
    y = _textbox(p1, pymupdf.Rect(72, y, 72 + width, y + 120), " ".join(INTRO))
    _textbox(p1, pymupdf.Rect(72, y, 72 + width, 800), " ".join(FILLER))

    p2 = doc.new_page()
    y = _heading(p2, 50, "2 Photosynthesis Experiments")
    y = _textbox(p2, pymupdf.Rect(72, y, 72 + width, y + 160), " ".join(PHOTO))
    _table(p2, y, YIELD_TABLE)

    p3 = doc.new_page()
    y = _heading(p3, 50, "3 Orbital Mechanics")
    _textbox(p3, pymupdf.Rect(72, y, 72 + width, y + 200), " ".join(ORBIT))
    doc.save(path)
    doc.close()
    return path


def make_scanned_pdf(path: Path) -> Path:
    """A single image-only page (what a scanner produces)."""
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 400, 300), False)
    pix.set_rect(pix.irect, (240, 240, 240))
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_image(pymupdf.Rect(72, 72, 472, 372), pixmap=pix)
    doc.save(path)
    doc.close()
    return path


def make_png(path: Path) -> Path:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 64, 64), False)
    pix.set_rect(pix.irect, (255, 255, 255))
    pix.save(str(path))
    return path


def make_docx(path: Path) -> Path:
    d = docx.Document()
    d.add_heading("Quarterly Review", level=0)
    d.add_heading("Revenue Analysis", level=1)
    for s in DOCX_REVENUE:
        d.add_paragraph(s)
    table = d.add_table(rows=3, cols=2)
    for r, (a, b) in enumerate([("Region", "Revenue"), ("Asia Pacific", "12.4M"), ("Europe", "8.1M")]):
        table.cell(r, 0).text, table.cell(r, 1).text = a, b
    d.add_page_break()
    d.add_heading("Hiring Plan", level=1)
    for s in DOCX_HIRING:
        d.add_paragraph(s)
    d.save(str(path))
    return path


def csv_rows(n: int = 300) -> list[list[str]]:
    rows = []
    for i in range(1, n + 1):
        city = "Zanzibar" if i == 217 else _CITIES[i % len(_CITIES)]
        rows.append([str(i), city, f"{20 + (i * 7) % 15}.{i % 10}", str((i * 37) % 400)])
    return rows


def make_csv(path: Path, n: int = 300) -> Path:
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(CSV_HEADER)
        writer.writerows(csv_rows(n))
    return path
