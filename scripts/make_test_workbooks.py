"""Build the Excel test workbooks in samples/test_workbooks/ plus answer_key.json.

All data is fictional. The bank statements copy the column layouts and narration styles of three real Indian
bank exports (SBI, ICICI, Kotak) but use merchants, employers and people that do NOT appear in the categoriser's
training generator (backend/ml/categoriser_data.py), so the categoriser is tested on names it has never seen.
Every row carries the label we assigned while writing it; the answer key holds the exact totals and ratios.

Usage: backend/.venv/Scripts/python scripts/make_test_workbooks.py
"""
from __future__ import annotations

import json
import random
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "samples" / "test_workbooks"

UNSEEN = {
    "groceries": ["EASYDAY CLUB", "SPAR HYPERMARKET", "HERITAGE FRESH", "FOODHALL", "NAMDHARIS FRESH", "KPN FARM FRESH"],
    "dining": ["VIDYARTHI BHAVAN", "PARADISE BIRYANI", "KAILASH PARBAT", "MAD OVER DONUTS", "GOLI VADA PAV", "JUMBOKING"],
    "travel": ["INDRIVE", "QUICK RIDE", "EASEMYTRIP", "GSRTC", "PUNE METRO", "UTSONMOBILE"],
    "fuel": ["GULF PETROCHEM", "MRPL FUEL STATION", "SRI VENKATESHWARA FUELS", "INDUS PETRO"],
    "utilities": ["JVVNL", "MGVCL", "SPECTRA BROADBAND", "GTPL BROADBAND", "PUNE MUNICIPAL WATER"],
    "shopping": ["TATA NEU", "SHOPCLUES", "HEALTHKART", "CAMPUS SHOES", "CHUMBAK"],
    "subscriptions": ["JIOSAAVN PRO", "KUKU FM", "POCKET FM", "SHEMAROOME", "HOICHOI"],
    "emi_loan": ["SMFG INDIA CREDIT", "IIFL FINANCE", "AU SMALL FINANCE BANK", "UJJIVAN SFB"],
    "investments": ["FYERS SECURITIES", "SAMCO SECURITIES", "DSP MUTUAL FUND", "TATA MUTUAL FUND", "EDELWEISS MF"],
}
PEOPLE = ["ANKIT VERMA", "PRIYA NAIR", "KARAN SHAH", "NEHA IYER", "VIVEK RAO"]
CITIES = ["PUNE", "JAIPUR", "VADODARA", "KOCHI"]
# (category, monthly count, min, max) of debits; salary/rent/EMI/SIP are added once a month below
MIX = [("groceries", 7, 180, 2600), ("dining", 6, 120, 1400), ("travel", 5, 60, 2400), ("fuel", 2, 500, 2500),
       ("utilities", 3, 299, 2400), ("shopping", 3, 450, 5200), ("subscriptions", 2, 99, 399),
       ("transfers", 2, 500, 4000), ("cash", 1, 2000, 5000)]


@dataclass
class Tx:
    day: date
    label: str      # categoriser label (backend/ml/categoriser_data.py LABELS)
    party: str      # merchant / employer / person
    amount: float   # positive
    credit: bool


def month_txs(rng: random.Random, start: date, months: int, employer: str, salary: float, rent: float,
              landlord: str, lender: str, emi: float, amc: str, sip: float) -> list[Tx]:
    out: list[Tx] = []
    for m in range(months):
        first = date(start.year + (start.month - 1 + m) // 12, (start.month - 1 + m) % 12 + 1, 1)
        nxt = date(first.year + first.month // 12, first.month % 12 + 1, 1)
        days = (nxt - first).days
        out.append(Tx(first, "salary", employer, salary, True))
        out.append(Tx(first + timedelta(days=2), "rent", landlord, rent, False))
        out.append(Tx(first + timedelta(days=4), "emi_loan", lender, emi, False))
        out.append(Tx(first + timedelta(days=9), "investments", amc, sip, False))
        out.append(Tx(first + timedelta(days=rng.randrange(days)), "transfers", rng.choice(PEOPLE),
                      float(rng.randrange(500, 3000, 50)), True))
        out.append(Tx(first + timedelta(days=days - 2), "fees", "SERVICE CHARGES GST", 17.7, False))
        for label, n, lo, hi in MIX:
            for _ in range(n):
                party = rng.choice(PEOPLE) if label == "transfers" else rng.choice(CITIES) if label == "cash" \
                    else rng.choice(UNSEEN[label])
                amount = float(rng.randrange(lo, hi, 500)) if label == "cash" else round(rng.uniform(lo, hi), 2)
                out.append(Tx(first + timedelta(days=rng.randrange(days)), label, party, amount, False))
    out.sort(key=lambda t: (t.day, not t.credit))
    return out


def ref(rng: random.Random, n: int = 12) -> str:
    return "".join(rng.choice("0123456789") for _ in range(n))


def vpa(party: str) -> str:
    return party.lower().replace(" ", "")[:14] + "@" + ("ybl" if len(party) % 2 else "okaxis")


# ── narration styles ──────────────────────────────────────────────────────────
def sbi_narration(rng: random.Random, t: Tx) -> str:
    r = ref(rng)
    if t.label == "salary":
        return f"BY TRANSFER-NEFT*UTIB0000553*AXN{ref(rng, 9)}*{t.party} SALARY--"
    if t.label == "rent":
        return f"TO TRANSFER-IMPS/{r}/{t.party}/RENT {t.day:%b}--".upper()
    if t.label == "emi_loan":
        return f"DEBIT-ACHDr NACH{ref(rng, 8)} {t.party}"
    if t.label == "investments":
        return f"DEBIT-ACHDr BD{ref(rng, 8)} {t.party} SIP"
    if t.label == "cash":
        return f"ATM WDL-ATM CASH {ref(rng, 4)} {t.party} MAIN RD"
    if t.label == "fees":
        return f"{t.party} {t.day:%b%y}".upper()
    if t.label == "utilities":
        return f"TO TRANSFER-BILLPAY/{t.party}/{r}--"
    if t.label == "transfers":
        way = "BY TRANSFER-UPI/CR" if t.credit else "TO TRANSFER-UPI/DR"
        return f"{way}/{r}/{t.party}/SBIN/{vpa(t.party)}/UPI--"
    if t.label == "fuel" and rng.random() < 0.5:
        return f"DEBIT CARD POS {ref(rng, 6)} {t.party} {rng.choice(CITIES)}"
    return f"TO TRANSFER-UPI/DR/{r}/{t.party}/YESB/{vpa(t.party)}/UPI--"


def icici_narration(rng: random.Random, t: Tx) -> str:
    r = ref(rng)
    if t.label == "salary":
        return f"NEFT-HDFCH00{ref(rng, 8)}-{t.party}-SALARY {t.day:%b%y}".upper()
    if t.label == "rent":
        return f"MMT/IMPS/{r}/Rent {t.day:%b}/{t.party}/HDFC Bank"
    if t.label == "emi_loan":
        return f"ACH/{t.party}/{ref(rng, 10)}"
    if t.label == "investments":
        return f"ACH/{t.party} SIP/{ref(rng, 10)}"
    if t.label == "cash":
        return f"NFS/CASH WDL/{r}/{t.day:%d-%m-%y}/{t.party}"
    if t.label == "fees":
        return f"{t.party}-{t.day:%b%y}".upper()
    if t.label == "utilities":
        return f"BIL/ONL/{ref(rng, 9)}/{t.party}/{ref(rng, 10)}"
    if t.label == "transfers":
        note = "Sent" if not t.credit else "Received"
        return f"UPI/{r}/{note}/{vpa(t.party)}/State Bank Of I/ICI{ref(rng, 8)}"
    if t.label in ("shopping", "fuel") and rng.random() < 0.5:
        return f"PCD/{ref(rng, 4)}/{t.party}/{rng.choice(CITIES)}/{ref(rng, 6)}"
    return f"UPI/{r}/Pay to {t.party.title()}/{vpa(t.party)}/YES BANK LIM/AXL{ref(rng, 8)}"


def kotak_narration(rng: random.Random, t: Tx) -> str:
    r = ref(rng)
    if t.label == "salary":
        return f"NEFT {ref(rng, 10)} {t.party} SAL {t.day:%b %Y}".upper()
    if t.label == "rent":
        return f"IMPS-{r}-{t.party}-RENT"
    if t.label == "emi_loan":
        return f"NACH-DR-{t.party}-{ref(rng, 8)}"
    if t.label == "investments":
        return f"NACH-DR-{t.party}-SIP{ref(rng, 6)}"
    if t.label == "cash":
        return f"ATL/{ref(rng, 6)}/{t.party}"
    if t.label == "fees":
        return t.party
    if t.label == "utilities":
        return f"IB BILLPAY DR-{t.party}-{ref(rng, 8)}"
    if t.label == "transfers":
        return f"UPI/{t.party}/{r}/{'Received' if t.credit else 'Sent'}"
    if t.label in ("dining", "shopping") and rng.random() < 0.4:
        return f"PCI/{ref(rng, 4)}/{t.party}/{rng.choice(CITIES)}"
    return f"UPI/{t.party}/{r}/UPI"


def balances(txs: list[Tx], opening: float) -> list[float]:
    out, bal = [], opening
    for t in txs:
        bal = round(bal + (t.amount if t.credit else -t.amount), 2)
        out.append(bal)
    return out


def write_rows(ws, rows: list[list]) -> None:
    for row in rows:
        ws.append(row)


def bold(ws, row: int) -> None:
    for cell in ws[row]:
        cell.font = Font(bold=True)


# ── statements ────────────────────────────────────────────────────────────────
def sbi(rng: random.Random) -> tuple[str, list[Tx]]:
    txs = month_txs(rng, date(2026, 7, 1), 3, "NIMBUS SOFTLABS PVT LTD", 92000, 24000, "SUNITA DESHPANDE",
                    "SMFG INDIA CREDIT", 8450, "DSP MUTUAL FUND", 5000)
    bal = balances(txs, 64210.40)
    wb = Workbook()
    ws = wb.active
    ws.title = "Statement"
    write_rows(ws, [["Account Name", ":", "Mr. ROHAN MEHTA"], ["Address", ":", "FLAT 12, BANER ROAD, PUNE 411045"],
                    ["Account Number", ":", "XXXXXXX4821"], ["Branch", ":", "BANER"], ["IFSC Code", ":", "SBIN0011234"],
                    ["Statement From", ":", "01-07-2026 to 30-09-2026"], [],
                    ["Txn Date", "Value Date", "Description", "Ref No./Cheque No.", "Debit", "Credit", "Balance"]])
    bold(ws, 8)
    for t, b in zip(txs, bal):
        amt = f"{t.amount:,.2f}"
        ws.append([f"{t.day.day} {t.day:%b %Y}", f"{t.day.day} {t.day:%b %Y}", sbi_narration(rng, t), ref(rng, 6),
                   "" if t.credit else amt, amt if t.credit else "", f"{b:,.2f}"])
    ws.append([])
    ws.append(["**This is a computer generated statement and does not require a signature."])
    name = "sbi_style_statement_q2fy27.xlsx"
    wb.save(OUT / name)
    return name, txs


def icici(rng: random.Random) -> tuple[str, list[Tx]]:
    txs = month_txs(rng, date(2026, 8, 1), 2, "KESTREL DATA SYSTEMS", 68500, 16500, "MOHAN JOSHI",
                    "IIFL FINANCE", 5120, "TATA MUTUAL FUND", 3000)
    bal = balances(txs, 31875.00)
    wb = Workbook()
    ws = wb.active
    ws.title = "OpTransactionHistory"
    write_rows(ws, [[None, "DETAILED STATEMENT"], [], [None, "Account Number", "XXXXXXXX7730 ( INR ) - PRIYANKA S"],
                    [None, "Transaction Date from", "01/08/2026", "to", "30/09/2026"], [],
                    [None, "S No.", "Value Date", "Transaction Date", "Cheque Number", "Transaction Remarks",
                     "Withdrawal Amount (INR )", "Deposit Amount (INR )", "Balance (INR )"]])
    bold(ws, 6)
    for i, (t, b) in enumerate(zip(txs, bal), 1):
        d = t.day.strftime("%d/%m/%Y")
        ws.append([None, i, d, d, "-", icici_narration(rng, t), 0 if t.credit else t.amount,
                   t.amount if t.credit else 0, b])
    name = "icici_style_statement_aug_sep_2026.xlsx"
    wb.save(OUT / name)
    return name, txs


def kotak(rng: random.Random) -> tuple[str, list[Tx]]:
    txs = month_txs(rng, date(2026, 9, 1), 1, "ORBIT LEARNING FOUNDATION", 41000, 9500, "LALITHA AMMA",
                    "UJJIVAN SFB", 2310, "EDELWEISS MF", 1000)
    bal = balances(txs, 12640.75)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    write_rows(ws, [["Kotak Mahindra Bank"], ["Account Statement 01-09-2026 to 30-09-2026"],
                    ["Cust. Reln. No.", "XXXX9921"], [],
                    ["Sl. No.", "Date", "Narration", "Chq/Ref No", "Amount", "Dr / Cr", "Balance", "Dr / Cr"]])
    bold(ws, 5)
    for i, (t, b) in enumerate(zip(txs, bal), 1):
        ws.append([i, t.day.strftime("%d-%m-%Y"), kotak_narration(rng, t), ref(rng, 10), t.amount,
                   "CR" if t.credit else "DR", abs(b), "CR" if b >= 0 else "DR"])
    name = "kotak_style_statement_sep_2026.xlsx"
    wb.save(OUT / name)
    return name, txs


def statement_key(txs: list[Tx]) -> dict:
    by_cat: dict[str, float] = {}
    for t in txs:
        by_cat[t.label] = round(by_cat.get(t.label, 0.0) + t.amount, 2)
    inflow = round(sum(t.amount for t in txs if t.credit), 2)
    outflow = round(sum(t.amount for t in txs if not t.credit), 2)
    return {"kind": "bank_statement", "rows": len(txs), "inflow": inflow, "outflow": outflow,
            "net": round(inflow - outflow, 2), "period": [min(t.day for t in txs).isoformat(),
                                                          max(t.day for t in txs).isoformat()],
            "byCategory": dict(sorted(by_cat.items(), key=lambda kv: -kv[1])),
            "subscriptions": sorted({t.party for t in txs if t.label == "subscriptions"})}


# ── balance sheets ────────────────────────────────────────────────────────────
def company() -> tuple[str, dict]:
    """Schedule III balance sheet + P&L of a fictional listed company, ₹ crore, two years side by side."""
    bs = [  # particulars, FY26, FY25, kind
        ("EQUITY AND LIABILITIES", None, None, None),
        ("Shareholders' funds", None, None, None),
        ("Share capital", 120.00, 120.00, "equity"), ("Reserves and surplus", 1846.35, 1612.80, "equity"),
        ("Non-current liabilities", None, None, None),
        ("Long-term borrowings", 410.25, 455.00, "ncl"), ("Deferred tax liabilities (net)", 38.60, 34.10, "ncl"),
        ("Long-term provisions", 52.40, 47.75, "ncl"),
        ("Current liabilities", None, None, None),
        ("Short-term borrowings", 145.00, 190.50, "cl"), ("Trade payables", 512.80, 468.20, "cl"),
        ("Other current liabilities", 214.65, 198.30, "cl"), ("Short-term provisions", 66.10, 59.40, "cl"),
        ("TOTAL", None, None, "total"),
        ("ASSETS", None, None, None),
        ("Non-current assets", None, None, None),
        ("Property, plant and equipment", 1388.70, 1302.15, "nca"), ("Capital work-in-progress", 96.40, 141.25, "nca"),
        ("Non-current investments", 240.00, 210.00, "nca"), ("Long-term loans and advances", 61.85, 58.40, "nca"),
        ("Current assets", None, None, None),
        ("Current investments", 155.30, 98.60, "ca"), ("Inventories", 604.20, 571.85, "ca"),
        ("Trade receivables", 498.75, 455.10, "ca"), ("Cash and cash equivalents", 252.65, 249.10, "ca"),
        ("Short-term loans and advances", 74.30, 69.90, "ca"), ("Other current assets", 34.00, 29.70, "ca"),
        ("TOTAL", None, None, "total"),
    ]
    sums: dict[str, float] = {}
    for _, fy26, _, kind in bs:
        if fy26 is not None:
            sums[kind] = round(sums.get(kind, 0.0) + fy26, 2)
    total_eq_liab = round(sums["equity"] + sums["ncl"] + sums["cl"], 2)
    total_assets = round(sums["nca"] + sums["ca"], 2)
    assert abs(total_eq_liab - total_assets) < 0.01, (total_eq_liab, total_assets)
    wb = Workbook()
    ws = wb.active
    ws.title = "Balance Sheet"
    write_rows(ws, [["SAHYADRI AGRO FOODS LIMITED (fictional)"], ["Standalone Balance Sheet as at 31 March 2026"],
                    ["(All amounts in ₹ crore, unless otherwise stated)"], [],
                    ["Particulars", "Note No.", "As at 31 March 2026", "As at 31 March 2025"]])
    bold(ws, 5)
    total25 = round(sum(fy25 for _, _, fy25, kind in bs if fy25 is not None and kind in ("nca", "ca")), 2)
    note = 2
    for p, fy26, fy25, kind in bs:
        if kind == "total":
            ws.append([p, None, total_assets, total25])
        elif fy26 is None:
            ws.append([p])
        else:
            note += 1
            ws.append([p, str(note), fy26, fy25])
    pl = wb.create_sheet("Profit and Loss")
    write_rows(pl, [["Statement of Profit and Loss for the year ended 31 March 2026 (₹ crore)"], [],
                    ["Particulars", "FY 2025-26", "FY 2024-25"],
                    ["Revenue from operations", 4215.60, 3870.25], ["Other income", 38.40, 31.05],
                    ["Total income", 4254.00, 3901.30],
                    ["Cost of materials consumed", 2690.15, 2501.40], ["Employee benefits expense", 402.30, 371.65],
                    ["Finance costs", 61.75, 68.90], ["Depreciation and amortisation expense", 128.40, 117.20],
                    ["Other expenses", 607.85, 560.15], ["Total expenses", 3890.45, 3619.30],
                    ["Profit before tax", 363.55, 282.00], ["Tax expense", 91.60, 71.10],
                    ["Profit for the year", 271.95, 210.90]])
    bold(pl, 3)
    name = "company_financials_fy26.xlsx"
    wb.save(OUT / name)
    equity = sums["equity"]
    liabilities = round(sums["ncl"] + sums["cl"], 2)
    return name, {"kind": "balance_sheet", "unit": "INR crore", "column": "As at 31 March 2026",
                  "totalAssets": total_assets, "totalLiabilities": liabilities, "equity": equity,
                  "currentRatio": round(sums["ca"] / sums["cl"], 3),
                  "quickRatio": round((sums["ca"] - 604.20) / sums["cl"], 3),
                  "debtToEquity_liabilities": round(liabilities / equity, 3),
                  "debtToEquity_borrowings": round((410.25 + 145.00) / equity, 3),
                  "netProfitMargin_pct": round(271.95 / 4215.60 * 100, 2),
                  "note": "Two year columns: the analyser should use FY26 (the first). The P&L is on sheet 2."}


def msme() -> tuple[str, dict]:
    """Meera's bakery: a small-business balance sheet the way an accountant's Excel usually looks."""
    rows = [("Capital Account", 850000, "equity"), ("Add: Net profit for the year", 312400, "equity"),
            ("Less: Drawings", -180000, "equity"),
            ("Term loan - Canara Bank (oven & equipment)", 420000, "ncl"),
            ("Sundry creditors (flour, dairy suppliers)", 168500, "cl"), ("GST payable", 23400, "cl"),
            ("Outstanding salaries", 54000, "cl"), ("Bank overdraft", 75000, "cl"),
            ("Shop premises (leasehold improvements)", 640000, "nca"), ("Ovens and kitchen equipment", 510000, "nca"),
            ("Delivery scooter", 85000, "nca"),
            ("Closing stock", 96500, "ca"), ("Sundry debtors (corporate orders)", 118200, "ca"),
            ("Cash in hand", 22600, "ca"), ("Cash at bank", 206000, "ca"), ("Prepaid rent", 45000, "ca")]
    sums: dict[str, float] = {}
    for _, v, k in rows:
        sums[k] = sums.get(k, 0) + v
    assert sums["equity"] + sums["ncl"] + sums["cl"] == sums["nca"] + sums["ca"], sums
    wb = Workbook()
    ws = wb.active
    ws.title = "BS 31-03-2026"
    write_rows(ws, [["MEERA'S BAKE HOUSE (Proprietor: Meera Kulkarni) — fictional"],
                    ["Balance Sheet as on 31st March 2026"], [], ["Particulars", "Amount (Rs.)"],
                    ["LIABILITIES", None]])
    bold(ws, 4)
    for p, v, k in rows:
        if k == "nca" and p.startswith("Shop"):
            ws.append(["Total Liabilities", sums["equity"] + sums["ncl"] + sums["cl"]])
            ws.append([])
            ws.append(["ASSETS", None])
        ws.append([p, v])
    ws.append(["Total Assets", sums["nca"] + sums["ca"]])
    name = "msme_bakery_balance_sheet_fy26.xlsx"
    wb.save(OUT / name)
    liabilities = sums["ncl"] + sums["cl"]
    return name, {"kind": "balance_sheet", "unit": "INR", "totalAssets": sums["nca"] + sums["ca"],
                  "totalLiabilities": liabilities, "equity": sums["equity"],
                  "currentRatio": round(sums["ca"] / sums["cl"], 3),
                  "quickRatio": round((sums["ca"] - 96500 - 45000) / sums["cl"], 3),
                  "debtToEquity_liabilities": round(liabilities / sums["equity"], 3),
                  "note": "Proprietorship layout: capital + profit - drawings is equity; overdraft is current."}


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    rng = random.Random(20261022)
    key: dict = {"_about": "Fictional test workbooks; labels and totals written by the Sphura team. Bank layouts "
                           "mimic SBI / ICICI / Kotak exports; merchants are absent from the categoriser's "
                           "training generator.", "files": {}}
    for build in (sbi, icici, kotak):
        name, txs = build(rng)
        key["files"][name] = statement_key(txs)
        key["files"][name]["labels"] = [[t.day.isoformat(), t.label, t.party, t.amount, "credit" if t.credit else
                                         "debit"] for t in txs]
    for build in (company, msme):
        name, k = build()
        key["files"][name] = k
    (OUT / "answer_key.json").write_text(json.dumps(key, indent=1), encoding="utf-8")
    for name, k in key["files"].items():
        print(f"{name:45s} {k['kind']:15s} " + (f"rows={k['rows']} in={k['inflow']:,.0f} out={k['outflow']:,.0f}"
                                                 if k["kind"] == "bank_statement" else
                                                 f"CR={k['currentRatio']} QR={k['quickRatio']}"))


if __name__ == "__main__":
    main()
